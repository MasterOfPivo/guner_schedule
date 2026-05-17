import os
import re
import subprocess
import json
import openpyxl
from models import db, ParsedSchedule

# 🐾 Папка с расписаниями
SCHEDULES_DIR = 'schedules'

def parse_pdf_schedule(filepath):
    try:
        result = subprocess.run(['pdftotext', '-layout', filepath, '-'], capture_output=True, text=True, check=True)
        lines = result.stdout.splitlines()
    except Exception as e:
        print(f"Ошибка при чтении PDF: {e}")
        return None, None
        
    groups_line_idx = -1
    for i, line in enumerate(lines):
        if "Дни" in line and "Число" in line:
            groups_line_idx = i + 1
            break
            
    if groups_line_idx == -1 or groups_line_idx >= len(lines):
        return None, None
        
    group_line = lines[groups_line_idx]
    
    groups = []
    for match in re.finditer(r'\S+', group_line):
        name = match.group()
        start = match.start()
        end = match.end()
        center = (start + end) / 2
        groups.append({'name': name, 'center': center, 'start': start, 'end': end})
        
    for i in range(len(groups)):
        if i == 0:
            groups[i]['col_start'] = 30
        else:
            groups[i]['col_start'] = int((groups[i-1]['center'] + groups[i]['center']) / 2)
            
        if i == len(groups) - 1:
            groups[i]['col_end'] = 1000
        else:
            groups[i]['col_end'] = int((groups[i]['center'] + groups[i+1]['center']) / 2)

    schedule_data = []
    current_day = None
    current_date = None
    current_cells = {g['name']: [] for g in groups}
    
    day_pattern = re.compile(r'^(Понедельник|Вторник|Среда|Четверг|Пятница|Суббота|Воскресенье)\s+(\d+\s+[а-яА-Я]+)')
    
    for line in lines[groups_line_idx+1:]:
        line_stripped = line.strip()
        if not line_stripped:
            continue
            
        match = day_pattern.search(line)
        if match:
            if current_day:
                row = {'day': current_day, 'time': current_date, 'cells_info': {}}
                for i, g in enumerate(groups):
                    val = "\n".join(current_cells[g['name']]).strip()
                    row['cells_info'][i] = {'value': val, 'merged_min': i, 'merged_max': i}
                schedule_data.append(row)
                
            current_day = match.group(1)
            current_date = match.group(2)
            current_cells = {g['name']: [] for g in groups}
            line = " " * match.end() + line[match.end():]
        
        if current_day is None:
            continue
            
        for g in groups:
            start = g['col_start']
            end = g['col_end']
            if start < len(line):
                cell_text = line[start:end].strip()
                if cell_text and "ВЫХОДНОЙ" not in cell_text.upper():
                    current_cells[g['name']].append(cell_text)

    if current_day:
        row = {'day': current_day, 'time': current_date, 'cells_info': {}}
        for i, g in enumerate(groups):
            val = "\n".join(current_cells[g['name']]).strip()
            row['cells_info'][i] = {'value': val, 'merged_min': i, 'merged_max': i}
        schedule_data.append(row)
        
    group_mapping = {g['name']: [i] for i, g in enumerate(groups)}
    
    return schedule_data, group_mapping

def parse_xlsx_schedule(filepath, filename):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        
        # 🐾 НАСТРОЙКИ СТРОК: 
        if 'экзамен' in filename.lower():
            GROUP_ROW = 4
            DATA_START_ROW = 5
        else:
            GROUP_ROW = 3       # На какой строке в Excel написаны названия групп
            DATA_START_ROW = 4  # С какой строки начинаются сами пары
        
        # 1. Быстрый поиск объединенных ячеек
        merged_lookup = {}
        for merged_range in ws.merged_cells.ranges:
            top_left_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
            val = top_left_cell.value
            for r in range(merged_range.min_row, merged_range.max_row + 1):
                for c in range(merged_range.min_col, merged_range.max_col + 1):
                    merged_lookup[(r, c)] = {
                        'value': val,
                        'min_col': merged_range.min_col,
                        'max_col': merged_range.max_col
                    }

        # 2. Определяем группы (ищем на строке GROUP_ROW, начиная с 3 колонки)
        group_mapping = {}
        last_main_group = None
        
        for col_idx in range(3, ws.max_column + 1):
            if (GROUP_ROW, col_idx) in merged_lookup:
                cell_val = merged_lookup[(GROUP_ROW, col_idx)]['value']
            else:
                cell_val = ws.cell(row=GROUP_ROW, column=col_idx).value
                
            if cell_val:
                g_name = str(cell_val).strip()
                # Убираем "25 чел" чтобы не было дублей групп
                g_name = re.sub(r'\s+\d+\s*чел\.?$', '', g_name, flags=re.IGNORECASE).strip()
                
                last_main_group = g_name
                if last_main_group not in group_mapping:
                    group_mapping[last_main_group] = []
                group_mapping[last_main_group].append(col_idx)
            elif last_main_group:
                # Если пустая ячейка рядом с группой, значит это её подгруппа
                group_mapping[last_main_group].append(col_idx)

        # 3. Собираем расписание
        schedule_data = []
        current_day = ""
        
        # Читаем данные, начиная со строки DATA_START_ROW
        for row_idx in range(DATA_START_ROW, ws.max_row + 1):
            # Время (колонка 2)
            if (row_idx, 2) in merged_lookup:
                time_val = merged_lookup[(row_idx, 2)]['value']
            else:
                time_val = ws.cell(row=row_idx, column=2).value
                
            if not time_val:
                continue # Пропускаем пустые строки
                
            # День (колонка 1)
            if (row_idx, 1) in merged_lookup:
                day_val = merged_lookup[(row_idx, 1)]['value']
            else:
                day_val = ws.cell(row=row_idx, column=1).value
                
            if day_val:
                current_day = str(day_val).strip()
                
            # 🐾 Превращаем время в строку и отрезаем секунды, если они есть (08:00:00 -> 08:00)
            time_str = str(time_val).strip()
            if time_str.count(':') == 2: 
                time_str = ":".join(time_str.split(':')[:2])
                
            if 'экзамен' in filename.lower():
                month_fixes = {
                    'янв': 'января', 'фев': 'февраля', 'мар': 'марта', 'апр': 'апреля',
                    'май': 'мая', 'июн': 'июня', 'июл': 'июля', 'авг': 'августа',
                    'сен': 'сентября', 'окт': 'октября', 'ноя': 'ноября', 'дек': 'декабря'
                }
                for short_m, full_m in month_fixes.items():
                    if time_str.lower().endswith(short_m):
                        time_str = time_str[:-len(short_m)] + full_m
                        break
                
            row_dict = {
                'day': current_day,
                'time': time_str,
                'cells_info': {}
            }
            
            # Сохраняем инфу по каждой колонке групп
            for g_cols in group_mapping.values():
                for c in g_cols:
                    if (row_idx, c) in merged_lookup:
                        val = merged_lookup[(row_idx, c)]['value']
                        m_min = merged_lookup[(row_idx, c)]['min_col']
                        m_max = merged_lookup[(row_idx, c)]['max_col']
                    else:
                        val = ws.cell(row=row_idx, column=c).value
                        m_min = c
                        m_max = c
                        
                    clean_val = str(val).strip() if val is not None else ""
                    if clean_val.lower() == 'none' or "ВЫХОДНОЙ" in clean_val.upper().replace(" ", ""):
                        clean_val = ""
                    elif 'экзамен' in filename.lower() and clean_val:
                        clean_val = re.sub(r'\s+', ' ', clean_val).strip()
                        
                    row_dict['cells_info'][c] = {
                        'value': clean_val,
                        'merged_min': m_min,
                        'merged_max': m_max
                    }
            
            schedule_data.append(row_dict)
            
        return schedule_data, group_mapping
    except Exception as e:
        print(f"Мяу-ошибка при чтении файла: {e}")
        return None, None

def get_processed_data(filename):
    filepath = os.path.join(SCHEDULES_DIR, filename)
    if not os.path.exists(filepath):
        return None, None
        
    current_mtime = os.path.getmtime(filepath)
    
    # 🐾 Проверяем кэш в БД
    record = ParsedSchedule.query.filter_by(filename=filename).first()
    if record and record.mtime == current_mtime:
        data = json.loads(record.data_json)
        g_map = json.loads(record.group_mapping_json)
        # Восстанавливаем ключи-числа (json делает их строками)
        for r in data:
            if 'cells_info' in r:
                r['cells_info'] = {int(k): v for k, v in r['cells_info'].items()}
        return data, g_map

    # Если в базе нет или файл обновился - парсим
    data = None
    g_map = None
    
    if filename.lower().endswith('.pdf'):
        data, g_map = parse_pdf_schedule(filepath)
    else:
        data, g_map = parse_xlsx_schedule(filepath, filename)
        
    if data is not None and g_map is not None:
        if record:
            record.mtime = current_mtime
            record.data_json = json.dumps(data, ensure_ascii=False)
            record.group_mapping_json = json.dumps(g_map, ensure_ascii=False)
        else:
            new_record = ParsedSchedule(
                filename=filename,
                mtime=current_mtime,
                data_json=json.dumps(data, ensure_ascii=False),
                group_mapping_json=json.dumps(g_map, ensure_ascii=False)
            )
            db.session.add(new_record)
        db.session.commit()

    return data, g_map

def get_all_schedules_info():
    if not os.path.exists(SCHEDULES_DIR):
        return [], []
    files = [f for f in os.listdir(SCHEDULES_DIR) if (f.endswith('.xlsx') or f.endswith('.pdf')) and not f.startswith('~$')]
    def sort_key(f):
        date_m = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', f)
        if date_m:
            day, month, year = date_m.groups()
            return (int(year), int(month), int(day))
        num_m = re.search(r'\d+', f)
        return (0, 0, int(num_m.group())) if num_m else (0, 0, 0)
        
    files.sort(key=sort_key, reverse=True)
    all_groups = set()
    for f in files:
        _, group_mapping = get_processed_data(f)
        if group_mapping:
            all_groups.update(group_mapping.keys())
            
    courses = {}
    other = []
    for g in sorted(all_groups):
        m = re.search(r'\d{2}', g)
        if m:
            year = int(m.group())
            courses.setdefault(year, []).append(g)
        else:
            other.append(g)
            
    sorted_courses = sorted(courses.keys(), reverse=True)
    structured_groups = []
    if sorted_courses:
        max_year = sorted_courses[0]
        for c in sorted_courses:
            course_num = max_year - c + 1
            structured_groups.append({
                'title': f"{course_num} курс",
                'groups': courses[c]
            })
    if other:
        structured_groups.append({
            'title': "Остальные",
            'groups': other
        })
        
    return files, structured_groups
